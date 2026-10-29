#!/usr/bin/env python3
"""Construye los datos del Mapa de Origen a partir del Excel de la red de proveedores.

Uso:  python3 scripts/construir_datos.py <Supply_Network.xlsx> <carpeta_geo> <carpeta_salida>

Entradas geográficas (carpeta_geo):
  ne_110m_admin_0_countries.geojson  (Natural Earth, países)
  COL_ADM1.geojson / COL_ADM2.geojson (geoBoundaries, departamentos y municipios de Colombia)

Salidas (carpeta_salida):
  mundo.json      países simplificados
  COL.json        departamentos + municipios (solo de los departamentos con proveedores)
  red.json        países → departamentos → municipios → proveedores → productos
"""
import json, sys, re, unicodedata, math
from collections import defaultdict, OrderedDict
import openpyxl

XLSX, GEO, OUT = sys.argv[1], sys.argv[2], sys.argv[3]

def norm(s):
    s = unicodedata.normalize("NFD", str(s or "").strip().lower())
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", s)

# ---------- geometría ----------
def perp(p, a, b):
    (x, y), (x1, y1), (x2, y2) = p, a, b
    dx, dy = x2 - x1, y2 - y1
    if dx == dy == 0:
        return math.hypot(x - x1, y - y1)
    t = max(0, min(1, ((x - x1) * dx + (y - y1) * dy) / (dx * dx + dy * dy)))
    return math.hypot(x - (x1 + t * dx), y - (y1 + t * dy))

def dp(pts, tol):
    if len(pts) < 3:
        return pts
    keep = [False] * len(pts); keep[0] = keep[-1] = True
    stack = [(0, len(pts) - 1)]
    while stack:
        i, j = stack.pop()
        dmax, idx = 0, None
        for k in range(i + 1, j):
            d = perp(pts[k], pts[i], pts[j])
            if d > dmax:
                dmax, idx = d, k
        if idx is not None and dmax > tol:
            keep[idx] = True
            stack += [(i, idx), (idx, j)]
    return [p for p, k in zip(pts, keep) if k]

def ring_area(r):
    return 0.5 * sum(r[i][0] * r[i + 1][1] - r[i + 1][0] * r[i][1] for i in range(len(r) - 1))

def polys(geom):
    return geom["coordinates"] if geom["type"] == "MultiPolygon" else [geom["coordinates"]]

def simplify(geom, tol, nd, min_area):
    out = []
    for poly in polys(geom):
        rings = []
        for i, r in enumerate(poly):
            if abs(ring_area(r)) < min_area and (i > 0 or len(polys(geom)) > 1):
                if i == 0:
                    break  # sin anillo exterior no hay polígono
                continue
            s = dp([tuple(p[:2]) for p in r], tol)
            s = [[round(x, nd), round(y, nd)] for x, y in s]
            dedup = [s[0]] + [p for k, p in enumerate(s[1:], 1) if p != s[k - 1]]
            if len(dedup) >= 4:
                # d3 (geometría esférica) exige anillo exterior en sentido horario y huecos antihorarios
                a = ring_area(dedup)
                if (len(rings) == 0 and a > 0) or (len(rings) > 0 and a < 0):
                    dedup = dedup[::-1]
                rings.append(dedup)
        if rings:
            out.append(rings)
    if not out:
        return None
    return {"type": "MultiPolygon", "coordinates": out} if len(out) > 1 else {"type": "Polygon", "coordinates": out[0]}

def centroid(geom):
    best = max((p[0] for p in polys(geom)), key=lambda r: abs(ring_area(r)))
    a = ring_area(best)
    if a == 0:
        return best[0][:2]
    cx = sum((best[i][0] + best[i + 1][0]) * (best[i][0] * best[i + 1][1] - best[i + 1][0] * best[i][1]) for i in range(len(best) - 1)) / (6 * a)
    cy = sum((best[i][1] + best[i + 1][1]) * (best[i][0] * best[i + 1][1] - best[i + 1][0] * best[i][1]) for i in range(len(best) - 1)) / (6 * a)
    return [cx, cy]

def in_ring(pt, r):
    x, y, inside = pt[0], pt[1], False
    for i in range(len(r) - 1):
        (x1, y1), (x2, y2) = r[i][:2], r[i + 1][:2]
        if (y1 > y) != (y2 > y) and x < (x2 - x1) * (y - y1) / (y2 - y1) + x1:
            inside = not inside
    return inside

def in_geom(pt, geom):
    return any(in_ring(pt, p[0]) and not any(in_ring(pt, h) for h in p[1:]) for p in polys(geom))

# ---------- Excel ----------
wb = openpyxl.load_workbook(XLSX, data_only=True)
prov_rows = [r for r in wb["Proveedores"].iter_rows(min_row=2, values_only=True) if r[0]]
prod_rows = [r for r in wb["Cropster Products"].iter_rows(min_row=2, values_only=True) if r[0] and r[1]]

# Correcciones de nombres de municipio: (municipio Excel, departamento) -> (municipio oficial, localidad, nota)
ALIAS = {
    ("arataco", "santander"): ("Aratoca", None, "En el Excel figura como «Arataco»."),
    ("dos quebradas", "risaralda"): ("Dosquebradas", None, None),
    ("ciudad bolivar", "antioquia"): ("Ciudad Bolívar", None, None),
    ("tablon de gomez", "narino"): ("El Tablón de Gómez", None, None),
    ("berruecos", "narino"): ("Arboleda", "Berruecos", "Berruecos es corregimiento de Arboleda."),
}

# ---------- límites ----------
adm1 = json.load(open(f"{GEO}/COL_ADM1.geojson"))["features"]
adm2 = json.load(open(f"{GEO}/COL_ADM2.geojson"))["features"]
DEP_FIX = {"Bogota Capital District": "Bogotá D.C.", "Archipiélago de San Andrés, Providencia y Santa Catalina": "San Andrés y Providencia"}
deps = []
for f in adm1:
    name = DEP_FIX.get(f["properties"]["shapeName"], f["properties"]["shapeName"])
    deps.append({"name": name, "iso": f["properties"]["shapeISO"], "geom": f["geometry"]})
dep_by_norm = {norm(d["name"]): d for d in deps}

munis = []
for f in adm2:
    c = centroid(f["geometry"])
    dep = next((d for d in deps if in_geom(c, d["geom"])), None)
    if dep is None:  # centroide fuera (costas): el departamento más cercano por centroide
        dep = min(deps, key=lambda d: math.dist(c, centroid(d["geom"])))
    name = f["properties"]["shapeName"]
    name = re.sub(r"\bDe\b", "de", name)
    munis.append({"name": name, "dep": dep["name"], "c": [round(c[0], 4), round(c[1], 4)], "geom": f["geometry"]})
muni_idx = {(norm(m["name"]), norm(m["dep"])): m for m in munis}

# ---------- red ----------
prods_by = defaultdict(list)
for pid, product, farm in prod_rows:
    prods_by[str(pid).strip()].append({"nombre": re.sub(r"\s+", " ", str(product).strip()), "finca": str(farm or "").strip()})

avisos = []
dep_tree = OrderedDict()
for r in prov_rows:
    pid, nombre, mun_raw, dep_raw = (str(x).strip() if x else "" for x in r[:4])
    dep = dep_by_norm.get(norm(dep_raw))
    if not dep:
        avisos.append(f"{pid} {nombre}: departamento «{dep_raw}» no reconocido."); continue
    key = (norm(mun_raw), norm(dep["name"]))
    mun_name, localidad, nota = ALIAS.get(key, (None, None, None))
    m = muni_idx.get((norm(mun_name), key[1])) if mun_name else muni_idx.get(key)
    if nota:
        avisos.append(f"{pid} {nombre}: {nota}")
    if m is None:
        avisos.append(f"{pid} {nombre}: no existe el municipio «{mun_raw}» en {dep['name']}; queda «por confirmar» en el centro del departamento.")
        mname, coord, ok = f"{mun_raw} (por confirmar)", [round(v, 4) for v in centroid(dep["geom"])], False
    else:
        mname, coord, ok = m["name"], m["c"], True
    d = dep_tree.setdefault(dep["name"], {"name": dep["name"], "iso": dep["iso"], "munis": OrderedDict()})
    mm = d["munis"].setdefault(mname, {"name": mname, "c": coord, "ok": ok, "proveedores": []})
    prods = prods_by.get(pid, [])
    if not prods:
        avisos.append(f"{pid} {nombre}: sin productos en «Cropster Products».")
    mm["proveedores"].append({"id": pid, "nombre": nombre, "localidad": localidad, "productos": prods})

ids = {str(r[0]).strip() for r in prov_rows}
for pid in prods_by:
    if pid not in ids:
        avisos.append(f"Productos con proveedor {pid}, que no está en la hoja «Proveedores».")

red = {
    "fuente": "Supply_Network.xlsx",
    "paises": [{
        "iso3": "COL", "nombre": "Colombia", "centro": [-73.3, 4.3],
        "departamentos": [{**{k: v for k, v in d.items() if k != "munis"}, "municipios": list(d["munis"].values())} for d in dep_tree.values()],
    }],
    "avisos": avisos,
}
json.dump(red, open(f"{OUT}/red.json", "w"), ensure_ascii=False, separators=(",", ":"))

# ---------- geo de salida ----------
active = set(dep_tree)
col = {"type": "FeatureCollection", "features": []}
for d in deps:
    g = simplify(d["geom"], 0.008, 3, 0.0005)
    if g:
        col["features"].append({"type": "Feature", "properties": {"t": "dep", "name": d["name"], "iso": d["iso"]}, "geometry": g})
for m in munis:
    if m["dep"] in active:
        g = simplify(m["geom"], 0.003, 3, 0.00005)
        if g:
            col["features"].append({"type": "Feature", "properties": {"t": "mun", "name": m["name"], "dep": m["dep"]}, "geometry": g})
json.dump(col, open(f"{OUT}/COL.json", "w"), ensure_ascii=False, separators=(",", ":"))

ne = json.load(open(f"{GEO}/ne_110m_admin_0_countries.geojson"))["features"]
mundo = {"type": "FeatureCollection", "features": []}
for f in ne:
    p = f["properties"]
    iso = p.get("ADM0_A3") or p.get("ISO_A3")
    g = simplify(f["geometry"], 0.05, 2, 0.01)
    if g:
        mundo["features"].append({"type": "Feature", "properties": {"iso3": iso, "name": p.get("NAME_ES") or p.get("NAME")}, "geometry": g})
json.dump(mundo, open(f"{OUT}/mundo.json", "w"), ensure_ascii=False, separators=(",", ":"))

print("Proveedores:", sum(len(m["proveedores"]) for d in dep_tree.values() for m in d["munis"].values()),
      "| Departamentos:", len(dep_tree), "| Municipios:", sum(len(d["munis"]) for d in dep_tree.values()),
      "| Productos:", sum(len(v) for v in prods_by.values()))
print("Avisos:"); [print(" -", a) for a in avisos]
